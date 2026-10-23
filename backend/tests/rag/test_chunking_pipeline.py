# backend/tests/rag/test_chunking_pipeline.py
from types import SimpleNamespace

import pytest

from backend.rag.preprocessing import metadata as _meta
from backend.rag.preprocessing.pipeline import parse_and_chunk

MD = """# 售后制度
## 退货流程
### 审核
客服审核退货原因。
## 差评处理
48小时内给出方案。
"""


@pytest.fixture
def deterministic_classify(monkeypatch):
    """fix f20：隔离 LLM 仲裁，避免单测依赖真实模型的非确定分类。

    售后制度 样本 policy(25)/sop(20) 胶着会触发仲裁：qwen 判 policy
    （Structure 策略，leaf+parent 双层），MiniMax 判 sop（Step 策略，
    仅 leaf）—— 两者语义都合理，单测必须锁定分类才可稳定断言。
    """
    monkeypatch.setattr(
        _meta, "llm",
        SimpleNamespace(invoke=lambda prompt: SimpleNamespace(content="policy")),
    )


def test_pipeline_end_to_end(tmp_path, deterministic_classify):
    p = tmp_path / "a.md"
    p.write_text(MD, encoding="utf-8")
    chunks = parse_and_chunk(str(p))
    assert chunks
    assert {"leaf", "parent"} <= {c.metadata["granularity"] for c in chunks}
    leaf = next(c for c in chunks if c.metadata["granularity"] == "leaf")
    assert leaf.metadata["parent_chunk_id"]
    assert leaf.metadata["section_path"]


def test_pipeline_unstructured_falls_back(tmp_path):
    t = tmp_path / "b.txt"
    t.write_text("客户提交申请后，客服核对订单信息。" * 100, encoding="utf-8")
    chunks = parse_and_chunk(str(t))
    assert chunks
    assert all(c.metadata.get("chunk_tokens", 0) > 0 for c in chunks)


def test_numbered_txt_paragraphs_not_lost(tmp_path):
    """编号 TXT 路由到 Structure 后段落不得被丢弃（TxtParser 段落需嵌套进 section）。"""
    t = tmp_path / "c.txt"
    t.write_text("一、退货流程\n提交申请。\n\n二、审核\n客服审核。\n", encoding="utf-8")
    chunks = parse_and_chunk(str(t))
    assert chunks
    all_text = "\n".join(c.page_content for c in chunks)
    assert "提交申请。" in all_text
    assert "客服审核。" in all_text


def test_document_title_fallback_to_filename(tmp_path):
    """解析器未识别标题时回退文件名；无语义文件名保持空（P1-1 回归）。

    PDF 等无标题样式格式的 document_title 曾全部为空，向量前缀的
    【文档】上下文与来源卡失去文档级锚点（2026-10-24）。
    """
    t = tmp_path / "员工手册_v2.txt"
    t.write_text("员工应当遵守公司各项规章制度。" * 40, encoding="utf-8")
    chunks = parse_and_chunk(str(t))
    titles = {c.metadata.get("document_title") for c in chunks}
    assert titles == {"员工手册 v2"}

    # 纯数字文件名不作为标题
    t2 = tmp_path / "2024.txt"
    t2.write_text("员工应当遵守公司各项规章制度。" * 40, encoding="utf-8")
    chunks2 = parse_and_chunk(str(t2))
    assert all(
        not (c.metadata.get("document_title") or "").strip() for c in chunks2
    )


def test_document_title_parser_wins_over_filename(tmp_path):
    """解析器显式识别的标题优先于文件名回退（DOCX Title 样式）。

    只有具备格式语义依据的解析器（DOCX 标题样式/HTML title/PPTX 首页
    占位符）设置 document_title；MD/TXT/PDF 无语义依据不设，走文件名回退。
    """
    import docx as _docx

    doc = _docx.Document()
    doc.add_heading("技术服务合同", level=1)
    doc.add_paragraph("第一条 服务内容：乙方向甲方提供平台运维服务。")
    doc.add_paragraph("第二条 服务费用：各项费用按季度结算。")
    p = tmp_path / "renamed_file.docx"
    doc.save(str(p))

    chunks = parse_and_chunk(str(p))
    titles = {c.metadata.get("document_title") for c in chunks}
    assert "技术服务合同" in titles
    assert "renamed file" not in titles
