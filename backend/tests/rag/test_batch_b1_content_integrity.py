from backend.rag.preprocessing.ast import DocumentAST, DocumentNode
from backend.rag.preprocessing.chunking import (
    FinancialTableChunkStrategy,
    RecursiveChunkStrategy,
    StructureChunkStrategy,
    coverage_audit,
)
from backend.rag.preprocessing.filter import ChunkFilter
from backend.rag.preprocessing.parser._table_nl import make_table_chunk_text


def test_pdf_cleaner_preserves_numeric_citation_lines():
    """法规文号中的年份和编号不能因版面断行被当作页码删除。"""
    from backend.rag.preprocessing.cleaner import DocumentCleaner

    cleaner = DocumentCleaner()
    citation = "依据文号：华师〔\n2024\n〕\n25\n号执行。"
    cleaned = cleaner.clean(citation, source_type="pdf").text

    assert "2024" in cleaned
    assert "25" in cleaned
    assert cleaner.clean("2024", source_type="pdf").text == "2024"
    assert cleaner.clean("123", source_type="pdf").text == "123"


def test_pdf_cleaner_preserves_ambiguous_numeric_content():
    """缺少页内坐标和跨页上下文时，保留可能属于正文的短数字。"""
    from backend.rag.preprocessing.cleaner import DocumentCleaner

    cleaner = DocumentCleaner()
    samples = (
        "依据华师〔\n2024\n〕25号执行。",
        "财务表格\n净利率\n2/5\n期末库存\n128",
        "产品手册\n额定电流\n16\n建议温度\n25",
        "123",
    )

    for sample in samples:
        assert cleaner.clean(sample, source_type="pdf").text == sample


def test_pdf_cleaner_removes_only_explicit_page_footer_lines(monkeypatch):
    """清理有明确标签的独立页脚，不误删相邻正文。"""
    from backend import config
    from backend.rag.preprocessing.cleaner import DocumentCleaner

    monkeypatch.setattr(config, "CLEAN_REMOVE_PDF_FOOTERS", True)
    cleaner = DocumentCleaner()
    result = cleaner.clean(
        "制度正文\nPage 2 of 10\n财务比率\n2/5\n第 3 页 共 10 页",
        source_type="pdf",
    )

    assert "Page 2 of 10" not in result.text
    assert "第 3 页 共 10 页" not in result.text
    assert "制度正文" in result.text
    assert "财务比率" in result.text
    assert "2/5" in result.text
    assert "removed_page_number_footer" in result.changes


def test_recursive_split_preserves_source_whitespace_and_coverage(monkeypatch):
    """递归切分不能删除边界空格或触发有效正文的 Coverage 硬失败。"""
    from backend.rag.preprocessing import chunking

    monkeypatch.setattr(chunking, "LEAF_CHUNK_TOKENS", 100)
    monkeypatch.setattr(chunking, "CHUNK_OVERLAP", 0)
    source = "甲" * 90 + " " + "乙" * 90
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="paragraph", text=source),
        ]),
        raw_text=source,
    )

    chunks = chunking.RecursiveChunkStrategy().split(ast, "whitespace.txt")
    leaves = [c for c in chunks if c.metadata.get("granularity") == "leaf"]

    assert len(leaves) > 1
    assert "".join(c.page_content for c in leaves) == source
    assert chunking.coverage_audit(ast, chunks)["uncovered_count"] == 0


def test_csv_table_routes_to_structured_chunks_and_preserves_identifiers(tmp_path):
    """CSV 表格的短数字行不可按普通中文正文规则过滤。"""
    from backend.rag.preprocessing.pipeline import parse_and_chunk_full

    path = tmp_path / "sku.csv"
    path.write_text(
        "SKU,订单号,金额\n"
        "000123,ORD-20261014-0000042,128.50\n"
        "000124,ORD-20261014-0000043,0.05\n",
        encoding="utf-8-sig",
    )
    chunks, qc = parse_and_chunk_full(str(path))
    assert qc["strategy_name"] == "StructureChunkStrategy"
    leaves = [c for c in chunks if c.metadata.get("granularity") == "leaf"]
    chunk_filter = ChunkFilter()
    kept = [
        c for c in leaves
        if chunk_filter.should_keep(c.page_content, c.metadata)[0]
    ]
    body = "\n".join(c.page_content for c in kept)
    assert "000123" in body
    assert "ORD-20261014-0000042" in body
    assert "128.50" in body
    assert len(kept) >= 2
    assert coverage_audit(qc["normalized_ast"], kept)["uncovered_count"] == 0
    ids = {c.metadata.get("chunk_id") for c in chunks}
    assert all(
        c.metadata.get("parent_chunk_id") in ids for c in kept
    )
    assert max(
        c.metadata.get("chunk_tokens", 0) for c in chunks
        if c.metadata.get("granularity") == "parent"
    ) <= 2000


def test_pdf_table_with_missing_cells_is_converted_without_none_crash():
    """PyMuPDF 表格中 None 单元格应成为空值，不能让 PDF 整体失败。"""
    result = make_table_chunk_text(
        [[None, "科目"], ["货币资金", None]],
        section_title="财务表",
    )
    assert "科目" in result
    assert "货币资金" in result
    assert "None" not in result


def test_excel_multilevel_header_rows_are_covered_as_leaf_content():
    """多级表头的原始每一行都需要进入可索引 Leaf。"""
    rows = [
        ["资产", "资产", "负债", "负债"],
        ["流动资产", "非流动资产", "流动负债", "非流动负债"],
        ["现金", "设备", "应付账款", "长期借款"],
        ["100.00", "200.00", "30.00", "40.00"],
    ]
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="section", text="财务报表", level=1, children=[
                DocumentNode(
                    type="table", text="财务表格", rows=rows,
                    children=[],
                ),
            ]),
        ]),
        raw_text="财务报表",
    )
    chunks = FinancialTableChunkStrategy().split(ast, "multi_sheet.xlsx")
    report = coverage_audit(ast, chunks)
    assert report["uncovered_count"] == 0, report
    leaves = [c for c in chunks if c.metadata.get("granularity") == "leaf"]
    assert any(
        "流动资产, 非流动资产, 流动负债, 非流动负债" in c.page_content
        for c in leaves
    )
    assert all(c.metadata.get("parent_chunk_id") for c in leaves)


def test_excel_formula_without_cached_value_keeps_formula_source(tmp_path):
    """公式没有缓存结果时，公式表达式也不能被当作空单元格丢弃。"""
    import openpyxl

    from backend.rag.preprocessing.parser.excel_parser import ExcelParser

    path = tmp_path / "formula.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "汇总"
    sheet.append(["收入", "成本", "结余"])
    sheet.append([42, 20, "=A2-B2"])
    workbook.save(path)
    workbook.close()

    ast = ExcelParser().parse(str(path))
    table = ast.root.children[0].children[0]

    assert table.rows[1] == ["42", "20", "=A2-B2"]


def test_excel_formula_uses_cached_result_when_available(tmp_path):
    """存在公式缓存值时优先索引缓存结果，不用公式文本覆盖。"""
    import xml.etree.ElementTree as ET
    from zipfile import ZIP_DEFLATED, ZipFile

    import openpyxl

    from backend.rag.preprocessing.parser.excel_parser import ExcelParser

    path = tmp_path / "cached_formula.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "汇总"
    sheet.append(["收入", "成本", "结余"])
    sheet.append([42, 20, "=A2-B2"])
    workbook.save(path)
    workbook.close()

    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    xml_path = "xl/worksheets/sheet1.xml"
    patched = tmp_path / "cached_formula_patched.xlsx"
    with ZipFile(path) as source, ZipFile(patched, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == xml_path:
                root = ET.fromstring(data)
                cell = root.find(f".//{{{namespace}}}c[@r='C2']")
                cached_value = cell.find(f"{{{namespace}}}v")
                cached_value.text = "22"
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(entry, data)

    ast = ExcelParser().parse(str(patched))
    table = ast.root.children[0].children[0]

    assert table.rows[1] == ["42", "20", "22"]


def test_coverage_audit_accepts_contiguous_leaf_chunk_boundaries():
    """连续 Leaf 的切块边界不能制造长段正文的 Coverage 缺口。"""
    text = "".join(
        f"第{i:04d}项操作须核验流程节点并记录处理结果。"
        for i in range(300)
    )
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="paragraph", text=text),
        ]),
        raw_text=text,
    )
    chunks = StructureChunkStrategy().split(ast, "long.txt")
    chunk_filter = ChunkFilter()
    kept = [
        chunk for chunk in chunks
        if chunk.metadata.get("granularity") != "leaf"
        or chunk_filter.should_keep(chunk.page_content, chunk.metadata)[0]
    ]
    split_leaves = [
        chunk for chunk in kept
        if chunk.metadata.get("granularity") == "leaf"
        and chunk.metadata.get("split_group_id")
    ]
    assert split_leaves
    assert {chunk.metadata.get("split_total") for chunk in split_leaves} == {
        len(split_leaves)
    }
    report = coverage_audit(ast, kept)
    assert report["uncovered_count"] == 0, report
    assert sum(
        c.metadata.get("granularity") == "leaf" for c in kept
    ) > 1


def test_coverage_joins_declared_parts_and_detects_a_missing_part():
    """边界重复内容不能误合并，分组缺片必须保留为覆盖缺口。"""
    from langchain_core.documents import Document

    motif = "".join(chr(0x4E00 + index) for index in range(60))
    left = "前" * 100 + motif
    right = motif + "后" * 100
    source = left + right
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="paragraph", text=source),
        ]),
        raw_text=source,
    )
    chunks = [
        Document(
            page_content=left,
            metadata={
                "granularity": "leaf", "split_group_id": "source-1",
                "split_part": 0, "split_total": 2,
            },
        ),
        Document(
            page_content=right,
            metadata={
                "granularity": "leaf", "split_group_id": "source-1",
                "split_part": 1, "split_total": 2,
            },
        ),
    ]

    complete = coverage_audit(ast, chunks)
    assert complete["uncovered_count"] == 0, complete

    incomplete = coverage_audit(ast, chunks[:1])
    assert incomplete["uncovered_count"] == 1, incomplete


def test_legal_long_clause_keeps_every_source_window():
    """长法律条款切分后，各输出片段仍须连续覆盖源文本。"""
    from backend.rag.preprocessing.chunking import LegalChunkStrategy

    source = "第一条\n" + "\n".join(
        f"第{i:04d}项教师工作要求应核验流程节点并记录处理结果。"
        for i in range(300)
    )
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="paragraph", text=source),
        ]),
        raw_text=source,
    )

    chunks = LegalChunkStrategy().split(ast, "long-legal.pdf")

    assert coverage_audit(ast, chunks)["uncovered_count"] == 0
    split_parts = sorted(
        (
            chunk.metadata["split_part"], chunk.page_content,
        )
        for chunk in chunks
        if chunk.metadata.get("granularity") == "leaf"
        and chunk.metadata.get("split_group_id")
    )
    assert "".join(text for _, text in split_parts) == source


def test_meaningful_english_pdf_text_is_not_rejected_by_chinese_ratio():
    """501 页英文 PDF 的真实英文正文不应仅因中文比例低而全部过滤。"""
    chunk_filter = ChunkFilter()
    text = (
        "The working tree stores each object by its content hash. "
        "Git can verify that the stored data remains unchanged. "
        "These commands describe the repository history."
    )
    kept, reason = chunk_filter.should_keep(
        text, {"doc_type": "general", "chunk_type": "paragraph"},
    )
    assert kept, reason


def test_pdf_source_identifier_survives_filter_and_coverage_audit():
    """PDF 页内有效编号不能因中文比例过滤而造成源叶覆盖缺口。"""
    from backend.rag.indexing.indexer import _filter_index_chunks

    source_text = "第1页 SENTINEL-OCR-P01"
    ast = DocumentAST(
        root=DocumentNode(
            type="section",
            text="",
            level=0,
            children=[DocumentNode(type="paragraph", text=source_text)],
        ),
        raw_text=source_text,
    )
    chunks = RecursiveChunkStrategy().split(ast, "source_text.pdf")
    filtered, filtered_details = _filter_index_chunks(
        chunks, ChunkFilter(), "source_text.pdf", "policy",
    )

    leaves = [
        chunk for chunk in filtered
        if chunk.metadata.get("granularity") == "leaf"
    ]
    assert any(source_text in chunk.page_content for chunk in leaves)
    assert not any(
        detail.get("reason") == "low_chinese_ratio"
        for detail in filtered_details
    )
    assert coverage_audit(ast, filtered)["uncovered_count"] == 0



def test_table_filter_preserves_short_and_duplicate_source_rows():
    """表格短行与重复行均是源记录；正文阈值和 SimHash 不能删除它们。"""
    chunk_filter = ChunkFilter()
    first = {"doc_type": "general", "chunk_type": "table_row"}
    second = {"doc_type": "general", "chunk_type": "table_row"}

    assert chunk_filter.should_keep("1", first) == (True, "clean")
    assert chunk_filter.should_keep("1", second) == (True, "clean")

    prose = {"doc_type": "general", "chunk_type": "paragraph"}
    assert chunk_filter.should_keep("1", prose) == (False, "too_short")



def test_coverage_treats_empty_pdf_table_cells_as_empty():
    """PyMuPDF 的 None 单元格是空白，不能生成虚假的缺失值。"""
    rows = [
        ["科目", "金额", "备注"],
        ["现金", "100.00", None],
        ["设备", None, "另表披露"],
    ]
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="table", text="", rows=rows),
        ]),
        raw_text="",
    )
    chunks = StructureChunkStrategy().split(ast, "empty-cells.pdf")
    report = coverage_audit(ast, chunks)
    assert report["uncovered_count"] == 0, report
    assert report["table_rows_uncovered"] == 0
    assert report["table_cells_uncovered"] == 0


def test_coverage_does_not_double_count_a_missing_duplicate_table_row():
    """Coverage 不能把同一个 Leaf 同时按原始行和拼接行重复计数。"""
    rows = [
        ["SKU", "商品", "价格"],
        ["000123", "商品A", "10.00"],
        ["000123", "商品A", "10.00"],
    ]
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="table", text="", rows=rows),
        ]),
        raw_text="",
    )
    chunks = StructureChunkStrategy().split(ast, "duplicate_rows.csv")
    retained = [
        chunk for chunk in chunks
        if not (
            chunk.metadata.get("chunk_type") == "table_row"
            and chunk.metadata.get("row_index") == 1
        )
    ]

    report = coverage_audit(ast, retained)

    assert report["table_rows_total"] == 3
    assert report["table_rows_uncovered"] == 1
    assert report["uncovered_count"] == 1


def test_coverage_falls_back_to_actual_leaf_for_single_row_table():
    """没有行级身份的单行 Sheet 不应被其他 Sheet 的行块误判为缺失。"""
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="section", text="员工", level=1, children=[
                DocumentNode(
                    type="table",
                    text=make_table_chunk_text(
                        [["工号", "姓名"], ["E0001", "员工1"]],
                        section_title="员工",
                    ),
                    rows=[["工号", "姓名"], ["E0001", "员工1"]],
                ),
            ]),
            DocumentNode(type="section", text="隐藏表", level=1, children=[
                DocumentNode(
                    type="table",
                    text=make_table_chunk_text(
                        [["secret", "SENTINEL-HIDDEN-SHEET"]],
                        section_title="隐藏表",
                    ),
                    rows=[["secret", "SENTINEL-HIDDEN-SHEET"]],
                ),
            ]),
        ]),
        raw_text="",
    )
    chunks = StructureChunkStrategy().split(ast, "multi_sheet.xlsx")

    report = coverage_audit(ast, chunks)

    assert report["uncovered_count"] == 0, report
    assert report["table_rows_uncovered"] == 0, report
    assert report["table_cells_uncovered"] == 0, report



def test_oversized_table_row_is_split_with_parent_and_full_coverage():
    """超长单元格分片仍遵守 Leaf 预算，并能按源行完成 Coverage。"""
    from backend.config import LEAF_CHUNK_TOKENS, PARENT_CHUNK_TOKENS

    rows = [
        ["商品", "说明"],
        ["长文档", "说明内容。" * 700, "10"],
    ]
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="table", text="", rows=rows),
        ]),
        raw_text="",
    )
    chunks = StructureChunkStrategy().split(ast, "wide-row.csv")
    leaves = [c for c in chunks if c.metadata.get("granularity") == "leaf"]
    parents = [c for c in chunks if c.metadata.get("granularity") == "parent"]
    ids = {c.metadata.get("chunk_id") for c in chunks}

    assert len(leaves) > 2
    assert all(c.metadata["chunk_tokens"] <= LEAF_CHUNK_TOKENS for c in leaves)
    assert all(c.metadata["chunk_tokens"] <= PARENT_CHUNK_TOKENS for c in parents)
    assert all(c.metadata.get("parent_chunk_id") in ids for c in leaves)
    row_body = "".join(
        c.page_content for c in leaves
        if c.metadata.get("chunk_type") == "table_row"
    )
    assert "说明内容。" * 700 in row_body
    report = coverage_audit(ast, chunks)
    assert report["uncovered_count"] == 0, report


def test_financial_table_section_keeps_bounded_paragraph_parents():
    """含表格 section 中的长正文也必须有实际 Parent 且遵守预算。"""
    from backend.config import LEAF_CHUNK_TOKENS, PARENT_CHUNK_TOKENS

    paragraph = DocumentNode(
        type="paragraph",
        text="流程依据应逐项核对并留存记录。" * 800,
    )
    table = DocumentNode(
        type="table", text="费用明细",
        rows=[["项目", "金额"], ["差旅费", "120.50"]],
        children=[],
    )
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(
                type="section", text="财务说明", level=1,
                children=[paragraph, table],
            ),
        ]),
        raw_text=paragraph.text,
    )

    chunks = FinancialTableChunkStrategy().split(ast, "financial.pdf")
    leaves = [c for c in chunks if c.metadata.get("granularity") == "leaf"]
    parents = [c for c in chunks if c.metadata.get("granularity") == "parent"]
    ids = {c.metadata.get("chunk_id") for c in chunks}

    assert leaves and parents
    assert all(c.metadata.get("parent_chunk_id") in ids for c in leaves)
    paragraph_body = "".join(
        c.page_content for c in leaves
        if not c.metadata.get("chunk_type", "").startswith("table_")
    )
    paragraph_leaves = [
        c for c in leaves
        if not c.metadata.get("chunk_type", "").startswith("table_")
    ]
    assert {c.metadata.get("split_total") for c in paragraph_leaves} == {
        len(paragraph_leaves)
    }
    assert paragraph.text in paragraph_body
    assert max(c.metadata["chunk_tokens"] for c in leaves) <= LEAF_CHUNK_TOKENS
    assert max(c.metadata["chunk_tokens"] for c in parents) <= PARENT_CHUNK_TOKENS
    report = coverage_audit(ast, chunks)
    assert report["uncovered_count"] == 0, report
