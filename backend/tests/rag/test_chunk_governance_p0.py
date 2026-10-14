"""test_chunk_governance_p0.py — RAG 切分链路 P0 正确性回归（2026-10-13）。

覆盖任务书 P0-01 ~ P0-05 的关键回归点：
- P0-01 混合 FAQ 不丢正文（parser 局部识别，见 test_qa_parser.py 补充用例）
- P0-02 AST 覆盖：根引导块 / 章节外叶子 / QA 混合文档 / 覆盖对账
- P0-03 Parent/Leaf ID 契约：索引期覆盖后 parent 关联可解析
- P0-04 表格按文档阅读顺序归属章节（DOCX/PDF）
- P0-05 统一 Token 上限：Step/Legal/QA/表格兜底/虚拟 parent
"""
import docx
import pytest

from backend.config import LEAF_CHUNK_TOKENS, PARENT_CHUNK_TOKENS
from backend.rag.preprocessing.ast import DocumentAST, DocumentNode, walk
from backend.rag.preprocessing.chunking import (
    ChunkStrategyRouter,
    QAChunkStrategy,
    StepChunkStrategy,
    StructureChunkStrategy,
    coverage_audit,
)
from backend.rag.preprocessing.parser.markdown_parser import MarkdownParser
from backend.rag.preprocessing.parser.txt_parser import TxtParser
from backend.rag.preprocessing.structure_analyzer import StructureAnalyzer
from backend.rag.preprocessing.token_counter import count_tokens


def _ast_of(md_text: str, tmp_path, suffix=".md") -> DocumentAST:
    p = tmp_path / f"doc{suffix}"
    p.write_text(md_text, encoding="utf-8")
    parser = MarkdownParser() if suffix == ".md" else TxtParser()
    return parser.parse(str(p))


# ── P0-02：AST 覆盖 ──────────────────────────────────────────────


def test_structure_covers_leaves_before_first_heading(tmp_path):
    """首个标题之前的正文（根直接子叶）不再被静默丢弃。"""
    ast = _ast_of(
        "这是标题前的引导段落。\n\n# 第一章\n\n第一章正文。\n", tmp_path)
    chunks = StructureChunkStrategy().split(ast, "doc.md")
    body = "\n".join(c.page_content for c in chunks)
    assert "引导段落" in body
    # 根引导块叶子挂在 parent 下
    leaf = next(
        c for c in chunks
        if "引导段落" in c.page_content
        and c.metadata["granularity"] == "leaf"
    )
    assert leaf.metadata["parent_chunk_id"]


def test_structure_skips_empty_container_sections(tmp_path):
    """纯容器 section（无直接叶子）不再产出空 parent。"""
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="section", text="总则", level=1, children=[
                DocumentNode(type="section", text="细则", level=2, children=[
                    DocumentNode(type="paragraph", text="细则内容。"),
                ]),
            ]),
        ]),
        raw_text="总则\n细则\n细则内容。",
    )
    chunks = StructureChunkStrategy().split(ast, "doc.md")
    # 「总则」无直接叶子 → 无空 parent；「细则」有叶子 → 正常产出
    assert not any(
        c.metadata.get("granularity") == "parent"
        and c.metadata.get("section_title") == "总则"
        and not c.page_content.strip("总则 \n")
        for c in chunks
    )
    assert any("细则内容" in c.page_content for c in chunks)


def test_step_strategy_covers_inter_section_leaves():
    """章节之间的根直接叶子（Step 策略）不再丢失。"""
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="paragraph", text="开篇说明。"),
            DocumentNode(type="section", text="一、准备", level=1, children=[
                DocumentNode(type="paragraph", text="准备步骤。"),
            ]),
            DocumentNode(type="paragraph", text="章节之间的插注。"),
            DocumentNode(type="section", text="二、执行", level=1, children=[
                DocumentNode(type="paragraph", text="执行步骤。"),
            ]),
        ]),
        raw_text="开篇说明。一、准备 准备步骤。章节之间的插注。二、执行 执行步骤。",
    )
    chunks = StepChunkStrategy().split(ast, "sop.txt")
    body = "\n".join(c.page_content for c in chunks)
    assert "开篇说明" in body
    assert "章节之间的插注" in body
    assert "执行步骤" in body


def test_qa_strategy_keeps_plain_paragraphs():
    """混合文档（QA + 普通段落）走 QA 策略时普通内容不丢。"""
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="qa_question", text="怎么退货？"),
            DocumentNode(type="qa_answer", text="提交申请。"),
            DocumentNode(type="paragraph", text="普通附则说明，不能丢。"),
        ]),
        raw_text="怎么退货？提交申请。普通附则说明，不能丢。",
    )
    chunks = QAChunkStrategy().split(ast, "mixed.md")
    body = "\n".join(c.page_content for c in chunks)
    assert "Q：怎么退货？" in body
    assert "普通附则说明" in body


def test_coverage_audit_detects_uncovered_leaf():
    """删除一个输出 Leaf 后，Parent 不得代替它通过覆盖对账。"""
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="section", text="A", level=1, children=[
                DocumentNode(type="paragraph", text="已覆盖的段落内容。"),
            ]),
            DocumentNode(type="section", text="B", level=1, children=[
                DocumentNode(type="paragraph", text="完全没进 chunk 的孤儿段落。"),
            ]),
        ]),
        raw_text="已覆盖的段落。完全没进 chunk 的孤儿段落。",
    )
    chunks = StructureChunkStrategy().split(ast, "doc.md")
    removed_leaf = next(
        c for c in chunks
        if c.metadata.get("granularity") == "leaf"
        and "孤儿段落" in c.page_content
    )
    chunks.remove(removed_leaf)

    report = coverage_audit(ast, chunks)
    assert report["uncovered_count"] == 1
    assert "孤儿段落" in report["uncovered_samples"][0]["preview"]
    assert report["covered"] == 1


def test_coverage_audit_full_coverage(tmp_path):
    """正常切分后覆盖率为 100%（叶子全部进 chunk）。"""
    ast = _ast_of(
        "# 章 A\n\n甲段落。\n\n## 小节\n\n乙段落。\n\n# 章 B\n\n丙段落。\n",
        tmp_path,
    )
    chunks = StructureChunkStrategy().split(ast, "doc.md")
    report = coverage_audit(ast, chunks)
    assert report["uncovered_count"] == 0, report


def test_coverage_audit_ignores_parent_chunks():
    """Parent 即使包含全文也不能替代检索实际输出的 Leaf。"""
    from langchain_core.documents import Document

    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="paragraph", text="这段正文必须进入 Leaf。"),
        ]),
        raw_text="这段正文必须进入 Leaf。",
    )
    parent_only = [Document(
        page_content="这段正文必须进入 Leaf。",
        metadata={"granularity": "parent"},
    )]

    report = coverage_audit(ast, parent_only)
    assert report["uncovered_count"] == 1


def test_coverage_audit_detects_missing_table_cell_and_row():
    """表格只有部分行/单元格进入 Leaf 时不能按单个命中判全表覆盖。"""
    from langchain_core.documents import Document

    rows = [
        ["SKU", "品名", "数量"],
        ["A001", "苹果", "100"],
        ["B002", "梨", "200"],
    ]
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="table", text="SKU 品名 数量 A001 苹果 100 B002 梨 200", rows=rows),
        ]),
        raw_text="SKU 品名 数量 A001 苹果 100 B002 梨 200",
    )
    chunks = [Document(
        page_content="SKU 品名 数量 A001 苹果 100 B002 梨",
        metadata={"granularity": "leaf"},
    )]

    report = coverage_audit(ast, chunks)
    assert report["uncovered_count"] == 1
    assert report["table_cells_uncovered"] >= 1


def test_sampled_coverage_never_claims_full_document_percentage():
    """抽样结果保留全量节点数与已审样本数，不输出全量 100%。"""
    from langchain_core.documents import Document

    texts = [f"文档段落{i}包含必须检索到的内容。" for i in range(8)]
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="paragraph", text=text) for text in texts
        ]),
        raw_text="\n".join(texts),
    )
    chunks = [Document(page_content=t, metadata={"granularity": "leaf"}) for t in texts]

    report = coverage_audit(ast, chunks, max_audit_nodes=2)
    assert report["sampled"] is True
    assert report["leaf_total"] == 8
    assert report["audited_leaf_count"] == 2
    assert report["coverage_rate"] is None


# ── P0-03：Parent/Leaf ID 契约 ───────────────────────────────────


def test_index_identity_overwrite_keeps_parent_linkage():
    """索引期 chunk_id 覆盖后，leaf.parent_chunk_id 指向 parent 的新 id。

    修复前：chunk_id 被 {doc_id}_{i} 覆盖而 parent_chunk_id 保留切分期
    内容派生 id → 检索端按 chunk_id $in 查 parent 永远落空。
    """
    from backend.rag.indexing.indexer import apply_index_chunk_identity

    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="qa_question", text="Q1？"),
            DocumentNode(type="qa_answer", text="A1。"),
        ]),
    )
    chunks = QAChunkStrategy().split(ast, "faq.md")
    split_parent_ids = {
        c.metadata["chunk_id"] for c in chunks
        if c.metadata["granularity"] == "parent"
    }
    apply_index_chunk_identity(chunks, "doc123")
    # 1) chunk_id 全部成为协议 id
    assert [c.metadata["chunk_id"] for c in chunks] == [
        f"doc123_{i}" for i in range(len(chunks))]
    # 2) 每个 leaf 的 parent_chunk_id 都能解析到真实 parent
    new_parent_ids = {
        c.metadata["chunk_id"] for c in chunks
        if c.metadata["granularity"] == "parent"
    }
    assert new_parent_ids.isdisjoint(split_parent_ids) or not split_parent_ids
    for c in chunks:
        if c.metadata["granularity"] == "leaf":
            assert c.metadata["parent_chunk_id"] in new_parent_ids


def test_parent_expansion_end_to_end_after_index_identity():
    """端到端：切分 → 索引身份覆盖 → 检索端 parent 扩展可命中。"""
    from backend.rag.indexing.indexer import apply_index_chunk_identity
    from backend.rag.retrieval.retrievers import attach_parent_context

    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="qa_question", text="Q1？"),
            DocumentNode(type="qa_answer", text="A1。"),
        ]),
    )
    chunks = QAChunkStrategy().split(ast, "faq.md")
    apply_index_chunk_identity(chunks, "docX")
    by_id = {c.metadata["chunk_id"]: c for c in chunks}

    def parent_lookup(ids):
        return [by_id[i] for i in ids if i in by_id]

    leaf = next(c for c in chunks if c.metadata["granularity"] == "leaf")
    result = attach_parent_context([leaf], parent_lookup)
    # parent 被成功拉入（修复前这里永远查不到）
    assert any(c.metadata["granularity"] == "parent" for c in result)


# ── P0-04：表格按阅读顺序 ────────────────────────────────────────


def _make_docx(path, blocks):
    """按顺序添加 heading/paragraph/table 块。"""
    d = docx.Document()
    for kind, payload in blocks:
        if kind == "h":
            d.add_heading(payload, level=1)
        elif kind == "p":
            d.add_paragraph(payload)
        else:
            rows, cols = payload
            t = d.add_table(rows=rows, cols=cols)
            for r in range(rows):
                for c in range(cols):
                    t.cell(r, c).text = f"r{r}c{c}"
    d.save(str(path))
    return str(path)


def test_docx_table_attaches_to_current_section(tmp_path):
    """正文与表格交错的 DOCX：表格归属阅读顺序中的当前章节。

    修复前：表格全部挂在最后一个 section。"""
    from backend.rag.preprocessing.parser.docx_parser import DocxParser

    p = tmp_path / "mixed.docx"
    path = _make_docx(p, [
        ("h", "第一章"),
        ("p", "第一章正文。"),
        ("t", (2, 2)),          # 第一章的表格
        ("h", "第二章"),
        ("p", "第二章正文。"),
        ("t", (2, 2)),          # 第二章的表格
    ])
    ast = DocxParser().parse(path)
    sections = {
        n.text: n for n in ast.root.children
        if n.type == "section" and n.level > 0
    }
    assert set(sections) == {"第一章", "第二章"}
    ch1_tables = [n for n in sections["第一章"].children if n.type == "table"]
    ch2_tables = [n for n in sections["第二章"].children if n.type == "table"]
    assert len(ch1_tables) == 1, "第一章的表格漂移到了其他章节"
    assert len(ch2_tables) == 1, "第二章的表格漂移到了其他章节"
    # 表格在各自章节的段落之后（阅读顺序保持）
    ch1_children = [n.type for n in sections["第一章"].children]
    assert ch1_children == ["paragraph", "table"]


def test_docx_mixed_faq_and_table_no_loss(tmp_path):
    """DOCX 含 FAQ + 表格 + 正文：三者都保留（修复前 FAQ 分支丢表格）。"""
    from backend.rag.preprocessing.parser.docx_parser import DocxParser

    p = tmp_path / "faq_table.docx"
    d = docx.Document()
    d.add_paragraph("Q：怎么退货？")
    d.add_paragraph("A：提交申请。")
    d.add_paragraph("普通说明段落。")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "SKU"
    t.cell(0, 1).text = "数量"
    t.cell(1, 0).text = "00123"
    t.cell(1, 1).text = "5"
    d.save(str(p))
    ast = DocxParser().parse(p)
    from backend.rag.preprocessing.ast import walk

    types = [n.type for n in walk(ast.root)]
    qa_nodes = [n for n in walk(ast.root) if n.type in ("qa_question", "qa_answer")]
    assert len(qa_nodes) == 2, "FAQ 对丢失"
    tables = [n for n in walk(ast.root) if n.type == "table"]
    assert len(tables) == 1, "表格丢失"
    assert any("00123" in c for row in tables[0].rows for c in row)
    # 普通段落内容不丢（紧随 FAQ 的正文按旧答案边界语义并入答案文本）
    all_text = "\n".join(n.text for n in walk(ast.root))
    assert "普通说明段落" in all_text


# ── P0-05：统一 Token 上限 ───────────────────────────────────────


def test_long_sop_step_split_preserves_identity():
    """超长 SOP 单步骤：切分后每片 ≤ Leaf 预算且保留章节身份。"""
    long_text = "这是超长步骤的详细说明。" * 400   # 远超 500 token
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="section", text="一、超长步骤", level=1, children=[
                DocumentNode(type="paragraph", text=long_text),
            ]),
        ]),
        raw_text=long_text,
    )
    chunks = StepChunkStrategy().split(ast, "sop.txt")
    leaves = [c for c in chunks if c.metadata["granularity"] == "leaf"]
    assert len(leaves) > 1, "超长步骤应被继续切分"
    for c in leaves:
        assert c.metadata["chunk_tokens"] <= LEAF_CHUNK_TOKENS
        assert c.metadata["section_title"] == "一、超长步骤"


def test_long_legal_clause_split_preserves_identity():
    """超长合同条款：切分后保留条款身份。"""
    clause_text = "甲方应当履行本条约定的全部义务，" * 300
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="paragraph", text="第一条 合同标的"),
            DocumentNode(type="paragraph", text=clause_text),
        ]),
        raw_text=clause_text,
    )
    from backend.rag.preprocessing.chunking import LegalChunkStrategy
    chunks = LegalChunkStrategy().split(ast, "contract.txt")
    leaves = [c for c in chunks if c.metadata["granularity"] == "leaf"]
    assert len(leaves) > 1
    for c in leaves:
        assert c.metadata["chunk_tokens"] <= LEAF_CHUNK_TOKENS
        assert "第一条" in c.metadata["section_title"]


def test_long_qa_answer_split_keeps_question_prefix():
    """FAQ 长答案：切分后每个子块保留 Q 前缀，且 ≤ Leaf 预算。"""
    long_answer = "答案的详细内容。" * 300
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="qa_question", text="怎么退货？"),
            DocumentNode(type="qa_answer", text=long_answer),
        ]),
        raw_text="怎么退货？" + long_answer,
    )
    chunks = QAChunkStrategy().split(ast, "faq.md")
    leaves = [c for c in chunks if c.metadata["granularity"] == "leaf"]
    assert len(leaves) > 1
    for c in leaves:
        assert c.page_content.startswith("Q：怎么退货？\nA：")
        assert c.metadata["chunk_tokens"] <= LEAF_CHUNK_TOKENS


def test_qa_virtual_parent_truncated_when_huge():
    """大量 QA 条目：虚拟 parent 超 PARENT 预算时截断并打标。"""
    children = []
    for i in range(300):
        children.append(DocumentNode(
            type="qa_question", text=f"问题{i}：" + "补充说明" * 30))
        children.append(DocumentNode(
            type="qa_answer", text=f"答案{i}。"))
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=children),
        raw_text="faq",
    )
    chunks = QAChunkStrategy().split(ast, "faq.md")
    parents = [c for c in chunks if c.metadata["granularity"] == "parent"]
    assert parents
    for p in parents:
        assert count_tokens(p.page_content) <= PARENT_CHUNK_TOKENS + 20
        assert p.metadata.get("parent_truncated") == "true"


def test_table_fallback_oversized_splits_by_lines():
    """表头规范化失败（仅一行、无数据行）的超大表兜底：不再产出单一巨型 chunk。"""
    from backend.rag.preprocessing.chunking import _split_table_node

    big_row = ["列" + "x" * 1000] * 8
    rows = [big_row]   # 仅表头行、无数据行 → normalize 失败走兜底
    node = DocumentNode(type="table", text="\n".join(",".join(r) for r in rows),
                        rows=rows)
    sec = DocumentNode(type="section", text="表章节", level=1)
    chunks = _split_table_node(node, sec, ["表章节"], "t.csv")
    leaves = [c for c in chunks if c.metadata["granularity"] == "leaf"]
    assert len(leaves) > 1, "超大兜底表应被切分"
    for c in leaves:
        assert c.metadata["chunk_tokens"] <= LEAF_CHUNK_TOKENS
        assert c.metadata["chunk_type"] == "table_fallback"
        # 兜底分片同样挂虚拟 parent（Parent 扩展不退化）
        assert c.metadata["parent_chunk_id"]


# ── 路由联合：混合 FAQ 文档经 Router 后内容完整 ─────────────────


def test_router_faq_full_pipeline_no_content_loss(tmp_path):
    """混合 FAQ 文档走 classify→router→split 全链路：QA 与正文都在。"""
    ast = _ast_of(
        "## 问题\n退货流程是什么？\n\n## 答案\n提交申请。\n\n"
        "# 第二章 附则\n\n附则正文不能丢。\n",
        tmp_path,
    )
    from backend.rag.preprocessing.metadata import classify_doc_type

    _, report = StructureAnalyzer().analyze(ast)
    doc_type = classify_doc_type(ast.raw_text, filename="faq.md")
    strategy = ChunkStrategyRouter().route(doc_type, report)
    chunks = strategy.split(ast, "faq.md")
    body = "\n".join(c.page_content for c in chunks)
    assert "退货流程" in body
    assert "附则正文" in body
