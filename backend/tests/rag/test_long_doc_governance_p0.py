"""test_long_doc_governance_p0.py — P0 超长文档完整处理回归（2026-10-13）。

覆盖任务书批次 2 验收点：
- CSV 前导零 SKU / 精确数字无损失（流式字符串读取不做数值推断）
- Excel 整数单元格无 ".0" 尾巴；行数上限截断留痕 AST
- chunk 超单文档上限 → 质量门禁硬异常（文档不得以截断状态 active）
- OCR 页数上限 / 失败页留痕（覆盖情况展示，不伪装完整成功）
"""
import csv
import io

import pytest

from backend.rag.preprocessing.ast import DocumentAST, DocumentNode
from backend.rag.preprocessing.parser.csv_parser import CsvParser
from backend.rag.preprocessing.parser.excel_parser import ExcelParser
from backend.rag.preprocessing.parser._table_nl import make_table_chunk_text
from backend.rag.preprocessing.quality_gate import (
    build_quality_record,
    hard_anomalies,
    run_typed_validation,
)


def _write_csv(tmp_path, name, header, rows):
    p = tmp_path / name
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    w.writerows(rows)
    p.write_text(buf.getvalue(), encoding="utf-8-sig")
    return str(p)


class TestCsvLeadingZero:
    def test_sku_leading_zeros_preserved(self, tmp_path):
        """前导零 SKU（"00123"）与长订单号不再被 pandas 数值推断破坏。"""
        path = _write_csv(
            tmp_path, "sku.csv",
            ["SKU", "订单号", "数量"],
            [["00123", "202610130001234567", "5"],
             ["00456", "202610130001234568", "12"]],
        )
        ast = CsvParser().parse(path)
        from backend.rag.preprocessing.ast import walk as _walk
        table = next(n for n in _walk(ast.root) if n.type == "table")
        flat = "\n".join(c for row in table.rows for c in row)
        assert "00123" in flat
        assert "00456" in flat
        assert "202610130001234567" in flat
        # 不再出现数值化的 "123" / 科学计数 / 浮点尾巴
        assert "00123.0" not in flat

    def test_gbk_csv_fallback_preserves_raw_strings(self, tmp_path):
        """GBK 回退仍保留中文、前导零与长订单号。"""
        path = tmp_path / "gbk.csv"
        path.write_bytes(
            "SKU,订单号,名称\n00123,202610130001234567,商品\n".encode("gbk")
        )
        ast = CsvParser().parse(str(path))
        from backend.rag.preprocessing.ast import walk as _walk

        table = next(n for n in _walk(ast.root) if n.type == "table")
        flat = "\n".join(cell for row in table.rows for cell in row)
        assert "00123" in flat
        assert "202610130001234567" in flat
        assert "商品" in flat

    def test_row_limit_truncation_marked(self, tmp_path, monkeypatch):
        import backend.config.rag as rag_cfg

        monkeypatch.setattr(rag_cfg, "RAG_TABLE_MAX_ROWS", 5)
        rows = [[f"sku{i}", str(i)] for i in range(20)]
        path = _write_csv(tmp_path, "big.csv", ["SKU", "Q"], rows)
        ast = CsvParser().parse(path)
        assert ast.table_rows_truncated > 0
        assert ast.table_rows_total == 20
        from backend.rag.preprocessing.ast import walk as _walk
        table = next(n for n in _walk(ast.root) if n.type == "table")
        assert len(table.rows) <= 5

    def test_csv_streaming_enforces_row_and_cell_caps_without_dataframe(self, tmp_path, monkeypatch):
        """超长和超宽 CSV 在读取时受限，且准确记录完整扫描得到的截断量。"""
        import pandas as pd
        import backend.config.rag as rag_cfg

        monkeypatch.setattr(rag_cfg, "RAG_TABLE_MAX_ROWS", 10_000)
        monkeypatch.setattr(rag_cfg, "RAG_TABLE_MAX_CELLS", 2_500)
        monkeypatch.setattr(
            pd, "read_csv",
            lambda *args, **kwargs: pytest.fail("CSV 不应整份载入 DataFrame"),
        )
        header = [f"列{i}" for i in range(1_000)]
        rows = [[f"值{r}-{c}" for c in range(1_000)] for r in range(100)]
        path = _write_csv(tmp_path, "超宽.csv", header, rows)

        ast = CsvParser().parse(path)
        from backend.rag.preprocessing.ast import walk as _walk

        table = next(n for n in _walk(ast.root) if n.type == "table")
        assert sum(map(len, table.rows)) <= 2_500
        assert ast.table_rows_total == 100
        assert ast.table_rows_truncated == 99
        assert ast.table_cells_total == 101_000
        assert ast.table_cells_truncated == 99_000


class TestExcelCellFidelity:
    def _xlsx(self, tmp_path, rows_by_sheet):
        import openpyxl

        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for sheet, rows in rows_by_sheet.items():
            ws = wb.create_sheet(sheet)
            for r in rows:
                ws.append(r)
        p = tmp_path / "t.xlsx"
        wb.save(str(p))
        return str(p)

    def test_integer_cells_have_no_float_tail(self, tmp_path):
        """数字单元格的 SKU（123456.0 形态）→ "123456"，无 .0 尾巴。"""
        path = self._xlsx(tmp_path, {"S1": [
            ["SKU", "数量"],
            [123456, 5.0],
            ["00123", "9"],   # 文本单元格原样保留（前导零）
        ]})
        ast = ExcelParser().parse(path)
        from backend.rag.preprocessing.ast import walk as _walk
        table = next(n for n in _walk(ast.root) if n.type == "table")
        flat = "\n".join(c for row in table.rows for c in row)
        assert "123456" in flat
        assert "123456.0" not in flat
        assert "00123" in flat
        assert ast.table_rows_truncated == 0

    def test_row_limit_truncation_marked(self, tmp_path, monkeypatch):
        import backend.config.rag as rag_cfg

        monkeypatch.setattr(rag_cfg, "RAG_TABLE_MAX_ROWS", 4)
        rows = [["h1", "h2"]] + [[f"v{i}", i] for i in range(30)]
        path = self._xlsx(tmp_path, {"S1": rows})
        ast = ExcelParser().parse(path)
        assert ast.table_rows_truncated > 0
        assert ast.table_rows_total == 4  # 不读取上限外的下一行
        assert ast.table_rows_total_known is False

    def test_workbook_limits_are_cumulative_and_stop_later_sheets(self, tmp_path, monkeypatch):
        """行/单元格上限按工作簿累计，触顶后不再读取后续 Sheet。"""
        import backend.config.rag as rag_cfg
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet

        monkeypatch.setattr(rag_cfg, "RAG_TABLE_MAX_ROWS", 5)
        monkeypatch.setattr(rag_cfg, "RAG_TABLE_MAX_CELLS", 100)
        original_iter_rows = ReadOnlyWorksheet.iter_rows
        visited_sheets = []

        def track_iter_rows(self, *args, **kwargs):
            visited_sheets.append(self.title)
            if self.title == "NeverRead":
                pytest.fail("达到工作簿上限后仍读取了后续 Sheet")
            yield from original_iter_rows(self, *args, **kwargs)

        monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", track_iter_rows)
        path = self._xlsx(tmp_path, {
            "First": [["ID", "值"], ["00123", 1], ["A2", 2]],
            "Second": [["ID", "值"], ["B1", 3], ["B2", 4]],
            "NeverRead": [["ID", "值"], ["C1", 5]],
        })

        ast = ExcelParser().parse(path)
        from backend.rag.preprocessing.ast import walk as _walk

        sections = [n.text for n in _walk(ast.root) if n.type == "section" and n.level == 1]
        tables = [n for n in _walk(ast.root) if n.type == "table"]
        assert sections == ["First", "Second"]
        assert set(visited_sheets) == {"First", "Second"}
        assert sum(len(t.rows) for t in tables) == 5
        assert ast.table_rows_truncated == 1
        assert ast.table_rows_total == 5  # 触顶即停，不读取第 6 行
        assert ast.table_rows_total_known is False
        assert "00123" in "\n".join(c for row in tables[0].rows for c in row)

    def test_cell_limit_is_cumulative_and_stops_before_next_sheet(
        self, tmp_path, monkeypatch,
    ):
        """单元格总预算跨 Sheet 累计，触顶后不实例化后续 Sheet。"""
        import backend.config.rag as rag_cfg
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet

        monkeypatch.setattr(rag_cfg, "RAG_TABLE_MAX_ROWS", 100)
        monkeypatch.setattr(rag_cfg, "RAG_TABLE_MAX_CELLS", 4)
        original_iter_rows = ReadOnlyWorksheet.iter_rows
        visited_sheets = []

        def track_iter_rows(self, *args, **kwargs):
            visited_sheets.append(self.title)
            if self.title == "NeverRead":
                pytest.fail("单元格达到工作簿上限后仍读取了后续 Sheet")
            yield from original_iter_rows(self, *args, **kwargs)

        monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", track_iter_rows)
        path = self._xlsx(tmp_path, {
            "First": [["h1", "h2"], ["a", "b"]],
            "NeverRead": [["h1", "h2"], ["c", "d"]],
        })

        ast = ExcelParser().parse(path)
        from backend.rag.preprocessing.ast import walk as _walk

        tables = [n for n in _walk(ast.root) if n.type == "table"]
        assert set(visited_sheets) == {"First"}
        assert sum(len(row) for table in tables for row in table.rows) == 4
        assert ast.table_cells_total == 4
        assert ast.table_cells_total_known is False
        assert ast.table_cells_truncated >= 1


class TestTruncatedDocNotActive:
    """chunk 超单文档上限 → 质量门禁硬异常（不得以截断状态 active）。"""

    def _record(self, tmp_path, truncated):
        ast = DocumentAST(
            root=DocumentNode(type="section", text="", level=0, children=[
                DocumentNode(type="paragraph", text="正文内容。"),
            ]),
            raw_text="正文内容。",
        )
        from langchain_core.documents import Document

        chunks = [Document(page_content="正文内容。", metadata={})]
        return build_quality_record(
            file_path=str(tmp_path / "doc.md"), doc_id="d1", kb_id="k1",
            raw_ast=ast, chunks=chunks, file_size=100, file_hash="h",
            truncated=truncated,
        )

    def test_truncation_is_hard_anomaly(self, tmp_path):
        record = self._record(tmp_path, truncated=True)
        anomalies = run_typed_validation(record, _empty_ast())
        hard = hard_anomalies(anomalies)
        assert any(a["check"] == "chunks_truncated" for a in hard), anomalies

    def test_no_truncation_no_hard_anomaly(self, tmp_path):
        record = self._record(tmp_path, truncated=False)
        anomalies = run_typed_validation(record, _empty_ast())
        hard = hard_anomalies(anomalies)
        assert not any(a["check"] == "chunks_truncated" for a in hard)


class TestOcrCoverageMarked:
    """OCR 页数上限 / 失败页留痕 → 不完整内容进入失败终态。"""

    def test_ocr_partial_coverage_flagged(self, tmp_path):
        ast = DocumentAST(
            root=DocumentNode(type="section", text="", level=0, children=[
                DocumentNode(type="paragraph", text="OCR 出来的内容。"),
            ]),
            raw_text="OCR 出来的内容。",
            ocr_required=True,
            ocr_attempted=True,
            ocr_triggered=True,
            ocr_pages=100,
            ocr_page_limit=100,
            ocr_page_budget=100,
            ocr_page_budget_source="default",
            ocr_failed_pages=3,
        )
        from langchain_core.documents import Document

        record = build_quality_record(
            file_path=str(tmp_path / "scan.pdf"), doc_id="d2", kb_id="k1",
            raw_ast=ast, chunks=[
                Document(page_content="OCR 出来的内容。", metadata={})],
            file_size=100, file_hash="h", ocr_triggered=True, ocr_pages=100,
        )
        anomalies = run_typed_validation(record, ast)
        partial = [a for a in anomalies if a["check"] == "ocr_coverage_partial"]
        assert partial, anomalies
        assert "页预算 100 页" in partial[0]["detail"]
        assert "上限遗漏 100 页" in partial[0]["detail"]
        assert "失败页 3" in partial[0]["detail"]
        # OCR 页数上限或失败页代表源内容缺失，不能以 partial 状态 active。
        assert partial[0]["severity"] == "error"


class TestTableTruncationFlagged:
    def test_table_rows_truncated_is_hard_error(self, tmp_path):
        ast = DocumentAST(
            root=DocumentNode(type="section", text="", level=0, children=[
                DocumentNode(type="table", text="h1,h2\nv1,1",
                             rows=[["h1", "h2"], ["v1", "1"]]),
            ]),
            raw_text="h1,h2",
            table_rows_truncated=900,
            table_rows_total=1000,
        )
        from langchain_core.documents import Document

        record = build_quality_record(
            file_path=str(tmp_path / "big.csv"), doc_id="d3", kb_id="k1",
            raw_ast=ast, chunks=[Document(page_content="h1,h2", metadata={})],
            file_size=100, file_hash="h",
        )
        anomalies = run_typed_validation(record, ast)
        flagged = [a for a in anomalies if a["check"] == "table_rows_truncated"]
        assert flagged
        assert "900" in flagged[0]["detail"]
        assert flagged[0]["severity"] == "error"


def test_financial_oversized_root_body_emits_bounded_linked_leaves(monkeypatch):
    """无标题超长财务正文要生成 Parent 与有链接的有界 Leaf。"""
    import backend.rag.preprocessing.chunking as chunking
    from backend.rag.preprocessing.chunking import FinancialTableChunkStrategy
    from backend.rag.preprocessing.ast import walk
    from backend.rag.preprocessing.chunking import count_tokens

    monkeypatch.setattr(chunking, "PARENT_CHUNK_TOKENS", 80)
    monkeypatch.setattr(chunking, "LEAF_CHUNK_TOKENS", 24)
    body = "\n".join(
        f"财务正文标记{i:04d}：本期经营收入和现金流数据持续披露。"
        for i in range(120)
    )
    ast = DocumentAST(
        root=DocumentNode(type="section", text="", level=0, children=[
            DocumentNode(type="paragraph", text=body),
        ]),
        source_file="无标题财务报告.txt",
        raw_text=body,
    )

    chunks = FinancialTableChunkStrategy().split(ast, ast.source_file)
    parents = [c for c in chunks if c.metadata.get("granularity") == "parent"]
    leaves = [c for c in chunks if c.metadata.get("granularity") == "leaf"]
    parent_ids = {c.metadata["chunk_id"] for c in parents}

    assert parents
    assert leaves
    assert all(c.metadata.get("parent_chunk_id") in parent_ids for c in leaves)
    assert all(count_tokens(c.page_content) <= 24 for c in leaves)
    assert all(count_tokens(c.page_content) <= 80 for c in parents)
    leaf_text = "\n".join(c.page_content for c in leaves)
    assert "财务正文标记0000" in leaf_text
    assert "财务正文标记0119" in leaf_text
    assert any(n.type == "paragraph" for n in walk(ast.root))


def _empty_ast() -> DocumentAST:
    return DocumentAST(
        root=DocumentNode(type="section", text="", level=0), raw_text="")
